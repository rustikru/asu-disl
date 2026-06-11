from __future__ import annotations

import imaplib
import json
import shutil
from datetime import datetime, timedelta
from email import message_from_bytes
from email.header import decode_header, make_header
from pathlib import Path

from werkzeug.utils import secure_filename

from .config import DEFAULT_SETTINGS, SETTINGS_FILE, TAB_APPROACH, TAB_DEPARTURE, get_user_data_file
from .models import MailFetchResult


def decode_mime_header(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (bytes, bytearray)):
        try:
            value = value.decode("utf-8")
        except Exception:
            try:
                value = value.decode("cp1251")
            except Exception:
                value = value.decode("latin-1", errors="ignore")
    text_value = str(value)
    if not text_value:
        return ""
    try:
        return str(make_header(decode_header(text_value)))
    except Exception:
        return text_value


def _message_id_text(message_id: object) -> str:
    if isinstance(message_id, (bytes, bytearray)):
        try:
            return message_id.decode("ascii", errors="ignore") or "mail"
        except Exception:
            return "mail"
    return str(message_id or "mail")


def load_settings(base_dir: Path) -> dict:
    path = get_user_data_file(SETTINGS_FILE)
    legacy_path = base_dir / SETTINGS_FILE
    if not path.exists() and legacy_path.exists():
        try:
            shutil.copy2(legacy_path, path)
        except Exception:
            pass
    if not path.exists():
        path.write_text(json.dumps(DEFAULT_SETTINGS, ensure_ascii=False, indent=2), encoding="utf-8")
        return json.loads(json.dumps(DEFAULT_SETTINGS))
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        merged = json.loads(json.dumps(DEFAULT_SETTINGS))
        legacy = data.get("imap", {}) if isinstance(data, dict) else {}
        merged["imap"] = {**DEFAULT_SETTINGS["imap"], **legacy}
        merged["imap_approach"] = {**DEFAULT_SETTINGS["imap_approach"], **legacy, **data.get("imap_approach", {})}
        merged["imap_departure"] = {**DEFAULT_SETTINGS["imap_departure"], **data.get("imap_departure", {})}
        return merged
    except Exception:
        return json.loads(json.dumps(DEFAULT_SETTINGS))


def get_imap_config(settings: dict, tab: str) -> dict:
    legacy = settings.get("imap", {}) if isinstance(settings, dict) else {}
    if tab == TAB_APPROACH:
        return {**DEFAULT_SETTINGS["imap_approach"], **legacy, **settings.get("imap_approach", {})}
    return {**DEFAULT_SETTINGS["imap_departure"], **settings.get("imap_departure", {})}


def remove_previous_mail_files(target_dir: Path, tab: str, keep_path: Path | None = None) -> None:
    if not target_dir.exists():
        return
    keep_resolved = None
    try:
        keep_resolved = keep_path.resolve() if keep_path is not None else None
    except Exception:
        keep_resolved = keep_path
    patterns = [f"{tab}_*.xlsx", f"{tab}_*.xls", f"{tab}_*.xlsm", f"{tab}_*.xltx", f"{tab}_*.xltm"]
    seen: set[Path] = set()
    for pattern in patterns:
        for path in target_dir.glob(pattern):
            if path in seen:
                continue
            seen.add(path)
            try:
                current = path.resolve()
            except Exception:
                current = path
            if keep_resolved is not None and current == keep_resolved:
                continue
            try:
                path.unlink()
            except Exception:
                pass


def fetch_latest_excel_from_mail(settings: dict, target_dir: Path, tab: str) -> MailFetchResult:
    imap_cfg = get_imap_config(settings, tab)
    if not imap_cfg.get("enabled"):
        raise ValueError("Получение почты для этой вкладки не включено в settings.json")
    server = imap_cfg.get("server", "").strip()
    username = imap_cfg.get("username", "").strip()
    password = imap_cfg.get("password", "")
    mailbox = imap_cfg.get("mailbox", "INBOX").strip() or "INBOX"
    sender_filter = imap_cfg.get("sender_filter", "").strip().lower()
    subject_filter = imap_cfg.get("subject_filter", "").strip().lower()
    subject_equals = imap_cfg.get("subject_equals", "").strip().lower()
    attachment_name_contains = imap_cfg.get("attachment_name_contains", "").strip().lower()
    attachment_name_equals = imap_cfg.get("attachment_name_equals", "").strip().lower()
    port = int(imap_cfg.get("port", 993) or 993)

    if not server or not username or not password:
        raise ValueError("В settings.json не заполнены server / username / password для IMAP.")

    target_dir.mkdir(parents=True, exist_ok=True)
    with imaplib.IMAP4_SSL(server, port) as client:
        client.login(username, password)
        status, _ = client.select(mailbox)
        if status != "OK":
            raise ValueError(f"Не удалось открыть папку {mailbox}.")
        recent_cutoff = datetime.now() - timedelta(days=1)
        since_date = recent_cutoff.strftime("%d-%b-%Y")
        status, data = client.search(None, "SINCE", since_date)
        if status != "OK":
            raise ValueError("Не удалось получить список свежих писем за последние сутки.")
        ids = data[0].split()
        if not ids:
            raise ValueError("За последние сутки в выбранной папке нет писем.")

        for message_id in reversed(ids):
            status, msg_data = client.fetch(message_id, "(RFC822)")
            if status != "OK":
                continue
            raw_email = msg_data[0][1]
            message = message_from_bytes(raw_email)
            from_value = decode_mime_header(message.get("From", ""))
            subject_value = decode_mime_header(message.get("Subject", ""))
            message_date = None
            try:
                from email.utils import parsedate_to_datetime
                if message.get("Date"):
                    message_date = parsedate_to_datetime(message.get("Date"))
                    if message_date.tzinfo is not None:
                        message_date = message_date.astimezone().replace(tzinfo=None)
            except Exception:
                message_date = None
            if message_date is not None and message_date < recent_cutoff:
                continue
            if sender_filter and sender_filter not in from_value.lower():
                continue
            subject_lower = subject_value.lower()
            if subject_equals and subject_lower != subject_equals:
                continue
            if subject_filter and subject_filter not in subject_lower:
                continue

            for part in message.walk():
                disposition = str(part.get("Content-Disposition", ""))
                if "attachment" not in disposition.lower():
                    continue
                filename = decode_mime_header(part.get_filename())
                if not filename:
                    continue
                lower_name = filename.lower()
                if not lower_name.endswith((".xlsx", ".xls", ".xlsm", ".xltx", ".xltm")):
                    continue
                if attachment_name_equals and lower_name != attachment_name_equals:
                    continue
                if attachment_name_contains and attachment_name_contains not in lower_name:
                    continue
                message_id_text = _message_id_text(message_id)
                safe_name = secure_filename(filename) or f"mail_attachment_{message_id_text}.xlsx"
                dated_prefix = datetime.now().strftime("%Y%m%d_%H%M%S")
                local_path = target_dir / f"{tab}_{dated_prefix}_{safe_name}"
                payload = part.get_payload(decode=True)
                if payload is None:
                    continue
                local_path.write_bytes(payload)
                remove_previous_mail_files(target_dir, tab, keep_path=local_path)
                return MailFetchResult(
                    file_path=str(local_path),
                    display_name=filename,
                    message_date=message_date,
                    from_value=from_value,
                    subject_value=subject_value,
                    signature=f"{message_id_text}::{filename}",
                )
    raise ValueError("Не найдено письмо с Excel-вложением по текущим фильтрам.")


def fetch_latest_excel_from_folder(target_dir: Path, tab: str, source_folder: Path = None) -> MailFetchResult:
    """Fetch latest Excel file from local folder instead of email."""
    if source_folder is None:
        # Try to find folder in multiple locations
        possible_paths = [
            Path(__file__).parent / tab,  # webapp/approach or webapp/departure
            Path.cwd() / tab,  # текущая директория
            Path(__file__).parent.parent / tab,  # Папка проекта (Python/approach)
            Path(__file__).resolve().parent / tab,  # Абсолютный путь webapp/approach
        ]
        source_folder = None
        for p in possible_paths:
            try:
                if p.exists() and p.is_dir():
                    source_folder = p
                    break
            except Exception:
                pass
        
        if source_folder is None:
            # Debug: print all tried paths
            debug_paths = "\n".join(str(p) for p in possible_paths)
            raise ValueError(f"Папка '{tab}' не найдена. Проверены пути:\n{debug_paths}")
    
    if not source_folder.exists():
        raise ValueError(f"Папка '{source_folder}' не существует.")
    
    # Find latest Excel file in folder
    excel_patterns = ["*.xlsx", "*.xls", "*.xlsm", "*.xltx", "*.xltm"]
    latest_file = None
    latest_mtime = 0
    
    for pattern in excel_patterns:
        for file_path in source_folder.glob(pattern):
            if file_path.is_file():
                mtime = file_path.stat().st_mtime
                if mtime > latest_mtime:
                    latest_mtime = mtime
                    latest_file = file_path
    
    if latest_file is None:
        raise ValueError(f"В папке '{source_folder}' не найдено Excel файлов.")
    
    # Copy file to target directory
    target_dir.mkdir(parents=True, exist_ok=True)
    dated_prefix = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = secure_filename(latest_file.name) or "data.xlsx"
    local_path = target_dir / f"{tab}_{dated_prefix}_{safe_name}"

    try:
        import openpyxl as _openpyxl
        _wb = _openpyxl.load_workbook(latest_file, read_only=True)
        _ws = _wb.active
        _first = [c.value for c in next(_ws.iter_rows(min_row=1, max_row=1))]
        _wb.close()
        # Если строка 1 содержит "Номер вагона" — это новый формат (без заголовка).
        # Нужно вставить 3 строки заголовка, как в формате почтового отчёта:
        #   Row 1: (пусто)
        #   Row 2: "DD.MM.YYYY HH:MM"  ← parse_report_datetime читает именно отсюда
        #   Row 3: "Данные о вагоне"
        #   Row 4: "Номер вагона", ...   ← заголовок столбцов (skiprows=3 в load_excel_as_df)
        _needs_header = any(v and "Номер вагона" in str(v) for v in _first)
    except Exception:
        _needs_header = False

    try:
        if _needs_header:
            import re as _re
            # Извлекаем дату из имени файла вида *_DDMMYYYY_HHMM.xlsx
            _dt_match = _re.search(r'(\d{2})(\d{2})(\d{4})_(\d{2})(\d{2})$', latest_file.stem)
            if _dt_match:
                _d, _m, _y, _h, _mi = _dt_match.groups()
                try:
                    _file_dt = datetime(int(_y), int(_m), int(_d), int(_h), int(_mi))
                except ValueError:
                    _file_dt = datetime.fromtimestamp(latest_mtime)
            else:
                _file_dt = datetime.fromtimestamp(latest_mtime)
            _date_str = _file_dt.strftime("%d.%m.%Y %H:%M")
            _wb = _openpyxl.load_workbook(latest_file)
            _ws = _wb.active
            _ws.insert_rows(1, amount=3)
            _ws["A2"] = _date_str  # дата в A2, как в почтовом формате
            _ws["A3"] = "Данные о вагоне"
            _wb.save(str(local_path))
        else:
            shutil.copy2(latest_file, local_path)
    except Exception as e:
        raise ValueError(f"Не удалось скопировать файл: {e}")
    
    remove_previous_mail_files(target_dir, tab, keep_path=local_path)
    
    return MailFetchResult(
        file_path=str(local_path),
        display_name=latest_file.name,
        message_date=datetime.now(),
        from_value="Локальная папка",
        subject_value=f"Файл из папки {tab}",
        signature=f"folder::{latest_file.name}::{int(latest_mtime)}",
    )