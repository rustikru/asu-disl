from __future__ import annotations

from typing import Iterable

import pandas as pd
from flask import current_app
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from core import strip_last_numeric_code
from core.parsers import parse_any_datetime


def _series_from_candidates(df: pd.DataFrame, candidates: Iterable[str], default: str = "") -> pd.Series:
    for col in candidates:
        if col in df.columns:
            series = df[col]
            if isinstance(series, pd.Series):
                return series
    return pd.Series([default] * len(df), index=df.index)


def _autosize_worksheet(ws, frame: pd.DataFrame) -> None:
    for idx, col in enumerate(frame.columns, start=1):
        values = [str(col)]
        if col in frame.columns:
            values.extend("" if pd.isna(x) else str(x) for x in frame[col].head(300).tolist())
        max_len = max((len(v) for v in values), default=12)
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = min(max(max_len + 2, 12), 40)
    ws.freeze_panes = "A2"




def _export_preferences() -> dict:
    try:
        prefs = dict(current_app.config.get("EXPORT_PREFERENCES") or {})
    except Exception:
        prefs = {}
    return {
        "excel_as_table": bool(prefs.get("excel_as_table", True)),
        "excel_with_filters": bool(prefs.get("excel_with_filters", True)),
        "excel_with_borders": bool(prefs.get("excel_with_borders", True)),
        "excel_freeze_header": bool(prefs.get("excel_freeze_header", True)),
        "excel_header_color": str(prefs.get("excel_header_color") or "dfeaff").replace("#", "")[:6] or "dfeaff",
        "excel_header_font_color": str(prefs.get("excel_header_font_color") or "1f3352").replace("#", "")[:6] or "1f3352",
    }


def _style_worksheet(ws, frame: pd.DataFrame) -> None:
    prefs = _export_preferences()
    _autosize_worksheet(ws, frame)
    if not prefs.get("excel_freeze_header", True):
        ws.freeze_panes = None
    header_fill = PatternFill("solid", fgColor=prefs.get("excel_header_color", "dfeaff"))
    header_font = Font(color=prefs.get("excel_header_font_color", "1f3352"), bold=True)
    thin = Side(style="thin", color="D6DEEA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if prefs.get("excel_as_table", True) and ws.max_row >= 1 and ws.max_column >= 1:
        ref = ws.dimensions
        if not ws.tables:
            table = Table(displayName=f"Table{abs(hash(ws.title)) % 100000}", ref=ref)
            style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            ws.add_table(table)
    if prefs.get("excel_with_filters", True):
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        if prefs.get("excel_with_borders", True):
            cell.border = border
    if prefs.get("excel_with_borders", True):
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = border

def _valid_departure_from_acceptance_series(df: pd.DataFrame) -> pd.Series:
    departure_series = _series_from_candidates(df, [
        "Дата и время отправления со станции приема",
        "Дата и время отправления со станции приёма",
        "Дата отправления со станции приема",
        "Дата отправления со станции приёма",
        "Дата и время отправления со станции приема груза",
        "Дата и время отправления со станции приёма груза",
    ])
    trip_start_series = _series_from_candidates(df, ["Дата и время начала рейса"])
    result = []
    for departure_value, trip_start_value in zip(departure_series.tolist(), trip_start_series.tolist()):
        departure_dt = parse_any_datetime(departure_value)
        trip_start_dt = parse_any_datetime(trip_start_value)
        if departure_dt is None or trip_start_dt is None or departure_dt < trip_start_dt:
            result.append("")
        else:
            result.append("" if pd.isna(departure_value) else str(departure_value))
    return pd.Series(result, index=df.index)

def detail_export_frame(df: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "№ ваг.",
        "Род ваг.",
        "Нач. рейса",
        "Дата и время окончания рейса",
        "Дор. отпр.",
        "Ст. отпр.",
        "Дор. назн.",
        "Ст. назн.",
        "Грузоотпр.",
        "Груз",
        "Ранее выгруженный груз",
        "Вес, кг",
        "Ст. опер.",
        "Дор. опер.",
        "Опер.",
        "Дата/время опер.",
        "Индекс поезда",
        "Контейнеры",
        "Норм. срок",
        "Пройд., км",
        "Ост., км",
        "Простой, сут.",
        "Отпр. со ст. пр.",
        "Приб. на ст. назн.",
        "Сост. ваг.",
        "Собств.",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    operation = _series_from_candidates(df, ["Операция с вагоном", "Операция"]).map(strip_last_numeric_code)

    return pd.DataFrame(
        {
            "№ ваг.": _series_from_candidates(df, ["Номер вагона"]),
            "Род ваг.": _series_from_candidates(df, ["Род вагона"]).map(strip_last_numeric_code),
            "Нач. рейса": _series_from_candidates(df, ["Дата и время начала рейса"]),
            "Дата и время окончания рейса": _series_from_candidates(df, ["Дата и время окончания рейса"]),
            "Дор. отпр.": _series_from_candidates(df, ["Дорога отправления"]),
            "Ст. отпр.": _series_from_candidates(df, ["Станция отправления"]),
            "Дор. назн.": _series_from_candidates(df, ["Дорога назначения"]),
            "Ст. назн.": _series_from_candidates(df, ["Станция назначения"]),
            "Грузоотпр.": _series_from_candidates(df, ["Грузоотправитель", "Грузоотправитель (наим)"]),
            "Груз": _series_from_candidates(df, ["Наименование груза"]),
            "Ранее выгруженный груз": _series_from_candidates(df, ["previous_cargo_display", "Ранее выгруженный груз"]),
            "Вес, кг": _series_from_candidates(df, ["Вес груза (кг)", "Вес груза"]),
            "Ст. опер.": _series_from_candidates(df, ["Станция операции"]),
            "Дор. опер.": _series_from_candidates(df, ["Дорога операции"]),
            "Опер.": operation,
            "Дата/время опер.": _series_from_candidates(df, ["Дата и время операции"]),
            "Индекс поезда": _series_from_candidates(df, ["Индекс поезда"]),
            "Контейнеры": _series_from_candidates(df, [
                "Номера контейнеров на вагоне",
                "Номера контейнеров",
                "Контейнеры",
                "№ контейнеров на вагоне",
            ]),
            "Норм. срок": _series_from_candidates(df, ["Нормативный срок доставки"]),
            "Пройд., км": _series_from_candidates(df, ["Расстояние пройденное (км)"]),
            "Ост., км": _series_from_candidates(df, ["Расстояние оставшееся (км)", "distance_left"]),
            "Простой, сут.": _series_from_candidates(df, ["Время простоя под последней операцией (сутки)"]),
            "Отпр. со ст. пр.": _valid_departure_from_acceptance_series(df),
            "Приб. на ст. назн.": _series_from_candidates(df, ["arrival_destination_display", "Дата и время прибытия (АСОУП) на станцию назначения"]),
            "Сост. ваг.": _series_from_candidates(df, ["Состояние вагона"]),
            "Собств.": _series_from_candidates(df, ["Собственник"]),
        }
    )


def manual_filter_export_frame(df: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "Номер вагона",
        "Род вагона",
        "Состояние вагона",
        "Собственник",
        "Наименование груза",
        "Грузоотправитель",
        "Грузополучатель",
        "Вес груза",
        "Станция отправления",
        "Дорога отправления",
        "Станция назначения",
        "Дорога назначения",
        "Станция операции",
        "Дорога операции",
        "Операция",
        "Индекс поезда",
        "Дата и время начала рейса",
        "Дата и время окончания рейса",
        "Дата и время операции",
    ]
    if df.empty:
        return pd.DataFrame(columns=columns)

    wagon_kind = _series_from_candidates(df, ["Род вагона"]).map(strip_last_numeric_code)
    operation = _series_from_candidates(df, ["Операция с вагоном", "Операция"]).map(strip_last_numeric_code)

    return pd.DataFrame(
        {
            "Номер вагона": _series_from_candidates(df, ["Номер вагона"]),
            "Род вагона": wagon_kind,
            "Состояние вагона": _series_from_candidates(df, ["wagon_state_display", "Состояние вагона"]),
            "Собственник": _series_from_candidates(df, ["owner_display", "Собственник"]),
            "Наименование груза": _series_from_candidates(df, ["cargo_name_display", "Наименование груза", "Груз"]),
            "Грузоотправитель": _series_from_candidates(df, ["Грузоотправитель"]),
            "Грузополучатель": _series_from_candidates(df, ["Грузополучатель"]),
            "Вес груза": _series_from_candidates(df, ["Вес груза (кг)", "Вес груза"]),
            "Станция отправления": _series_from_candidates(df, ["origin_station_display", "Станция отправления"]),
            "Дорога отправления": _series_from_candidates(df, ["origin_road_display", "Дорога отправления"]),
            "Станция назначения": _series_from_candidates(df, ["destination_display", "Станция назначения"]),
            "Дорога назначения": _series_from_candidates(df, ["destination_road_display", "Дорога назначения"]),
            "Станция операции": _series_from_candidates(df, ["station_display", "Станция операции"]),
            "Дорога операции": _series_from_candidates(df, ["road_display", "Дорога операции"]),
            "Операция": operation,
            "Индекс поезда": _series_from_candidates(df, ["Индекс поезда"]),
            "Дата и время начала рейса": _series_from_candidates(df, ["Дата и время начала рейса"]),
            "Дата и время окончания рейса": _series_from_candidates(df, ["Дата и время окончания рейса", "trip_end"]),
            "Дата и время операции": _series_from_candidates(df, ["Дата и время операции"]),
        }
    )


def write_excel_report(summary_df: pd.DataFrame, detail_df: pd.DataFrame, target_path: str) -> None:
    with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Сводка", index=False)
        detail_df.to_excel(writer, sheet_name="Пономерной список", index=False)

        for sheet_name, frame in {
            "Сводка": summary_df,
            "Пономерной список": detail_df,
        }.items():
            ws = writer.book[sheet_name]
            _style_worksheet(ws, frame)


def write_single_sheet_excel(
    frame: pd.DataFrame,
    target_path: str,
    sheet_name: str = "Результат",
) -> None:
    if frame is None:
        frame = pd.DataFrame()

    with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name=sheet_name[:31] or "Результат", index=False)
        ws = writer.book[sheet_name[:31] or "Результат"]
        _style_worksheet(ws, frame)